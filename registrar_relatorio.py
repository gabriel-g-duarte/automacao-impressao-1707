import os
import sys
from datetime import datetime
import json
import openpyxl

# Caminho do arquivo de configuração
CONFIG_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.json")

def carregar_pasta_projeto():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                cfg = json.load(f)
                return cfg.get("PASTA_PROJETO", os.path.dirname(os.path.abspath(__file__)))
        except Exception:
            pass
    return os.path.dirname(os.path.abspath(__file__))

def registrar_codigo_no_excel(codigo, status_impressao="Concluído"):
    pasta_destino = carregar_pasta_projeto()
    caminho_excel = os.path.join(pasta_destino, "relatorio_impressao.xlsx")

    # Data e Horário atual
    agora = datetime.now()
    data_str = agora.strftime("%d/%m/%Y")
    horario_str = agora.strftime("%H:%M:%S")

    # Garante que o código seja tratado como string limpa
    codigo_str = str(codigo).strip()

    # 1. Se a planilha não existe, cria do zero com os cabeçalhos
    if not os.path.exists(caminho_excel):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Relatório Impressão"
        ws.append(["Código", "Data", "Horário", "Status"])
    else:
        wb = openpyxl.load_workbook(caminho_excel)
        ws = wb.active

    # 2. Varre as linhas existentes para verificar se o código já existe
    encontrado = False
    for row in range(2, ws.max_row + 1):
        codigo_celula = str(ws.cell(row=row, column=1).value or "").strip()
        
        if codigo_celula == codigo_str:
            # Encontrou o código repetido: apenas atualiza o horário, data e status
            ws.cell(row=row, column=2).value = data_str
            ws.cell(row=row, column=3).value = horario_str
            ws.cell(row=row, column=4).value = status_impressao
            encontrado = True
            break

    # 3. Se não encontrou o código na planilha, adiciona uma nova linha
    if not encontrado:
        ws.append([codigo_str, data_str, horario_str, status_impressao])

    # Salva e fecha a planilha
    wb.save(caminho_excel)
    wb.close()

if __name__ == "__main__":
    # Permite receber o código via linha de comando ou teste direto
    if len(sys.argv) > 1:
        cod = sys.argv[1]
        registrar_codigo_no_excel(cod)
    else:
        print("Uso: python registrar_relatorio.py <codigo>")